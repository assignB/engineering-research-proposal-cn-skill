#!/usr/bin/env python3
"""Build the default DOCX template used when the user supplies no formal template."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("output", type=Path)
    return parser.parse_args()


def set_fonts(style, east_asia: str, western: str, size: float, bold: bool = False) -> None:
    style.font.name = western
    style.font.size = Pt(size)
    style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:ascii"), western)
    rfonts.set(qn("w:hAnsi"), western)
    rfonts.set(qn("w:cs"), western)


def set_run_fonts(run, east_asia: str, western: str, size: float, bold: bool = False) -> None:
    run.font.name = western
    run.font.size = Pt(size)
    run.font.bold = bold
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:ascii"), western)
    rfonts.set(qn("w:hAnsi"), western)


def add_field(paragraph, instruction: str, fallback: str = "") -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    paragraph.add_run()._r.append(begin)
    paragraph.add_run()._r.append(instr)
    paragraph.add_run()._r.append(separate)
    paragraph.add_run(fallback)
    paragraph.add_run()._r.append(end)


def add_numbering(doc: Document) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(e.get(qn("w:abstractNumId"))) for e in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(e.get(qn("w:numId"))) for e in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multilevel = OxmlElement("w:multiLevelType")
    multilevel.set(qn("w:val"), "multilevel")
    abstract.append(multilevel)
    for level in range(3):
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(level))
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), "decimal")
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), ".".join(f"%{i}" for i in range(1, level + 2)))
        suff = OxmlElement("w:suff")
        suff.set(qn("w:val"), "space")
        lvl.append(start)
        lvl.append(num_fmt)
        lvl.append(lvl_text)
        lvl.append(suff)
        abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_heading_numbering(doc: Document, num_id: int) -> None:
    for level in range(3):
        style = doc.styles[f"Heading {level + 1}"]
        ppr = style.element.get_or_add_pPr()
        numpr = ppr.find(qn("w:numPr"))
        if numpr is None:
            numpr = OxmlElement("w:numPr")
            ppr.append(numpr)
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), str(level))
        numid = OxmlElement("w:numId")
        numid.set(qn("w:val"), str(num_id))
        numpr.append(ilvl)
        numpr.append(numid)


def shade_cell(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    trpr.append(tbl_header)


def add_guidance(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="Template Guidance")
    p.add_run(f"【编写提示】{text}")


def add_metadata_table(doc: Document) -> None:
    table = doc.add_table(rows=5, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    rows = (
        ("项目名称", "待填写"),
        ("文档类型", "科研调研报告 / 技术论证报告 / 项目申报书"),
        ("编制单位", "待填写"),
        ("编制人员", "待填写"),
        ("编制日期", "待填写"),
    )
    for row, values in zip(table.rows, rows):
        row.height = Cm(0.9)
        for index, value in enumerate(values):
            cell = row.cells[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = value
            if index == 0:
                shade_cell(cell, "D9EAF2")
                for run in cell.paragraphs[0].runs:
                    run.bold = True


def add_control_table(doc: Document) -> None:
    headings = ["控制项", "当前设置", "确认状态"]
    rows = [
        ("责任范围", "整篇 / 指定章节", "待确认"),
        ("目标篇幅", "按项目配置", "待确认"),
        ("技术指标", "不在未确认时写死数值", "待确认"),
        ("格式模板", "用户模板优先；无模板时采用本文件", "已知"),
        ("公式对象", "MathType OLE 优先；实际类型必须披露", "待确认"),
        ("工程图源文件", "VSDX", "已知"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])
    for i, heading in enumerate(headings):
        table.rows[0].cells[i].text = heading
        shade_cell(table.rows[0].cells[i], "5B9BD5")
        for run in table.rows[0].cells[i].paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.bold = True
    for values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = value


def add_outline_section(doc: Document, title: str, guidance: str, subsections: tuple[str, ...]) -> None:
    doc.add_heading(title, level=1)
    add_guidance(doc, guidance)
    for subsection in subsections:
        doc.add_heading(subsection, level=2)
        add_guidance(doc, "围绕本标题形成完整论证单元；删除本提示后再形成正式正文。")


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    set_fonts(normal, "宋体", "Times New Roman", 12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.first_line_indent = Pt(24)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.widow_control = True

    for level, size in ((1, 15), (2, 14), (3, 12)):
        style = doc.styles[f"Heading {level}"]
        set_fonts(style, "黑体", "Times New Roman", size, True)
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.space_before = Pt(18 if level == 1 else 12)
        style.paragraph_format.space_after = Pt(8 if level == 1 else 6)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.page_break_before = level == 1
        style.element.get_or_add_pPr().append(OxmlElement("w:keepNext"))

    set_fonts(doc.styles["Title"], "黑体", "Times New Roman", 22, True)
    doc.styles["Title"].paragraph_format.first_line_indent = Pt(0)
    doc.styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_fonts(doc.styles["Subtitle"], "宋体", "Times New Roman", 14)
    doc.styles["Subtitle"].paragraph_format.first_line_indent = Pt(0)
    doc.styles["Subtitle"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_fonts(doc.styles["Caption"], "楷体", "Times New Roman", 10.5)
    doc.styles["Caption"].paragraph_format.first_line_indent = Pt(0)
    doc.styles["Caption"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    guidance = doc.styles.add_style("Template Guidance", 1)
    set_fonts(guidance, "楷体", "Times New Roman", 10.5)
    guidance.font.color.rgb = RGBColor(89, 89, 89)
    guidance.paragraph_format.first_line_indent = Pt(0)
    guidance.paragraph_format.line_spacing = 1.25
    guidance.paragraph_format.space_after = Pt(6)

    num_id = add_numbering(doc)
    apply_heading_numbering(doc, num_id)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.first_line_indent = Pt(0)
    add_field(footer, " PAGE ", "1")


def build(output: Path) -> None:
    doc = Document()
    configure_document(doc)
    props = doc.core_properties
    props.title = "中文工科科研调研、技术论证与项目申报书通用模板"
    props.subject = "未指定正式模板时使用的首版通用 Word 模板"
    props.author = "engineering-research-proposal-cn"

    for _ in range(5):
        doc.add_paragraph()
    title = doc.add_paragraph(style="Title")
    title.add_run("中文工科科研调研、技术论证与项目申报书")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run("通用写作模板（首版）")
    doc.add_paragraph()
    add_metadata_table(doc)
    note = doc.add_paragraph(style="Template Guidance")
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.add_run("本模板仅用于未提供正式规范时；获得正式模板后应按用户指定格式更新。")

    doc.add_page_break()
    doc.add_paragraph("编制说明", style="Title")
    add_guidance(doc, "先完成项目配置、责任边界和指标确认，再开始正文。团队协作时必须确认负责的标题等级、标题数量、既定编号及合稿接口。")
    add_control_table(doc)
    p = doc.add_paragraph()
    p.add_run("术语使用顺序：用户材料原词、现行标准与权威资料、本领域常用专业词汇。不得自行创造概念或缩略语。")
    p = doc.add_paragraph()
    p.add_run("同一上级标题下的并列小标题通常控制在 3—5 个；超过 5 个时优先重新分组。固定申报栏目除外。")

    doc.add_page_break()
    doc.add_paragraph("目  录", style="Title")
    toc = doc.add_paragraph()
    toc.paragraph_format.first_line_indent = Pt(0)
    add_field(toc, ' TOC \\o "1-3" \\h \\z \\u ', "打开 Word 后更新目录")

    outlines = (
        ("立项背景与需求分析", "说明任务来源、应用场景、工程需求和约束，不在未确认时写死指标。", ("应用背景与任务需求", "使用场景与约束条件", "拟解决的主要问题")),
        ("国内外研究现状与技术差距", "按技术问题组织真实文献和代表性工作，形成可追溯的比较结论。", ("国内研究现状", "国外研究现状", "现有方案比较与技术差距")),
        ("总体思路与技术路线", "给出从输入、处理到输出的总体关系，并说明各部分接口。", ("总体设计思路", "系统总体架构", "技术路线与实施路径")),
        ("研究内容与关键技术", "围绕核心技术对象展开；同级小标题一般控制在 3—5 个。", ("研究内容分解", "关键技术一", "关键技术二", "关键技术三")),
        ("系统方案设计与实现", "说明功能组成、软硬件划分、接口、数据流程和实现约束。", ("功能组成与模块划分", "接口关系与数据流程", "软硬件实现方案", "系统集成方法")),
        ("试验验证与考核方法", "将研究内容、指标来源、试验条件和判定方法对应起来。", ("验证对象与试验条件", "试验方法与数据处理", "指标判定与结果记录")),
        ("风险分析与保障措施", "从技术、进度、资源、接口和试验条件识别风险并提出可执行措施。", ("技术风险及应对措施", "进度与资源风险", "质量控制与组织保障")),
        ("预期成果与应用", "区分可交付成果、应用方式和后续扩展，不使用未经证明的宣传性判断。", ("预期成果形式", "应用场景与实施方式", "后续研究与工程化安排")),
    )
    for args in outlines:
        add_outline_section(doc, *args)

    doc.add_heading("参考文献", level=1)
    add_guidance(doc, "采用配置规定的著录格式；正文引用和文后条目必须一一对应，所有来源可核验。")
    doc.add_paragraph("[1] 作者. 题名[文献类型标识]. 出版信息, 年份.")

    doc.add_heading("附录：术语、来源与指标登记", level=1)
    doc.add_heading("术语与缩略语表", level=2)
    add_guidance(doc, "记录首选写法、英文全称、缩略语、来源和适用范围。")
    doc.add_heading("来源登记表", level=2)
    add_guidance(doc, "记录来源编号、题名、作者或机构、年份、链接或 DOI、适用章节和支持的判断。")
    doc.add_heading("技术指标登记表", level=2)
    add_guidance(doc, "记录指标名称、数值、单位、条件、来源、确认状态和适用章节；未确认值不得写成承诺。")

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(output.resolve())


def main() -> int:
    args = parse_args()
    build(args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
